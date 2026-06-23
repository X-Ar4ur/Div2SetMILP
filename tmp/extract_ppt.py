from pptx import Presentation
import io

path = r"D:\Work\UCASNJ\PPT\Div2SetMILP_v2.pptx"
prs = Presentation(path)
out = []
for i, slide in enumerate(prs.slides, 1):
    out.append("\n" + "="*70 + f"\n### SLIDE {i}\n" + "="*70)
    for shape in slide.shapes:
        if shape.has_text_frame:
            txt = shape.text_frame.text
            if txt.strip():
                out.append(f"[TEXT shape='{shape.name}']\n{txt}")
        if shape.has_table:
            tbl = shape.table
            out.append("[TABLE]")
            for r in tbl.rows:
                cells = [c.text for c in r.cells]
                out.append(" | ".join(cells))
        if shape.shape_type == 13:
            out.append(f"[PICTURE name='{shape.name}']")
    if slide.has_notes_slide:
        n = slide.notes_slide.notes_text_frame.text
        if n.strip():
            out.append(f"[NOTES]\n{n}")

with io.open(r"D:\Work\UCASNJ\EasyBC-2080-related-key\tmp\ppt_dump.txt","w",encoding="utf-8") as f:
    f.write("\n".join(out))
print("SLIDES:", len(prs.slides._sldIdLst))
