from pptx import Presentation
from pptx.util import Emu

path = r"D:\Work\UCASNJ\PPT\Div2SetMILP_v2.pptx"
prs = Presentation(path)
W, H = prs.slide_width, prs.slide_height
print(f"SLIDE SIZE EMU = {W} x {H}  (= {Emu(W).inches:.2f} x {Emu(H).inches:.2f} in)")

def emu_in(v):
    try:
        return f"{Emu(v).inches:.2f}"
    except Exception:
        return "?"

# Focus on slides that likely hold the iterativeSolver flowchart
for idx in [21, 26, 35]:
    slide = prs.slides[idx-1]
    print("\n" + "="*60 + f"\nSLIDE {idx}")
    for sh in slide.shapes:
        kind = sh.shape_type
        has_txt = sh.has_text_frame and sh.text_frame.text.strip()
        txt = ""
        if has_txt:
            txt = sh.text_frame.text.strip().replace("\n", " / ")[:40]
        L = emu_in(sh.left); T = emu_in(sh.top); Wd = emu_in(sh.width); Hh = emu_in(sh.height)
        print(f"  name='{sh.name}' type={kind} pos=({L},{T}) size=({Wd}x{Hh}) txt='{txt}'")
